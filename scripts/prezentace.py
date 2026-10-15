"""Prezentace KAM pro hackathon: velká písmena, tmavé pozadí, přechod Morph.

    python scripts/prezentace.py

Čísla bere z `data/kam.db` (stejně jako aplikace), snímky obrazovky
z `prezentace/img/` (`scripts/prezentace.py --snimky` je nafotí z běžícího
serveru na localhost:8000 headless Chromem). Výstup: `prezentace/KAM_Popojedem.pptx`.

Jak funguje pohyb: PowerPoint (2019 / 365) při přechodu Morph spáruje objekty,
které mají na sousedních snímcích **stejné jméno začínající „!!“**, a plynule
je přesune, zvětší, natočí a přebarví. Každý snímek proto nese tytéž objekty
(kruhy, nadpis, snímky aplikace, graf); co na snímku nemá být vidět, leží
mimo plochu a přiletí zvenku. Starší PowerPoint dostane místo Morphu prolnutí.
"""

from __future__ import annotations

import os
import sqlite3
import subprocess
import sys
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

KOREN = Path(__file__).resolve().parent.parent
VYSTUP = KOREN / "prezentace"
IMG = VYSTUP / "img"
CHROME = os.environ.get("CHROME", r"C:\Program Files\Google\Chrome\Application\chrome.exe")

POZADI = RGBColor(0x0D, 0x0F, 0x12)
BILA = RGBColor(0xF4, 0xF4, 0xF2)
SEDA = RGBColor(0x8B, 0x90, 0x99)
ZELENA = RGBColor(0x8E, 0xE0, 0x6A)
FIALOVA = RGBColor(0xA7, 0x8B, 0xFA)
MODRA = RGBColor(0x5A, 0xA9, 0xFF)
ORANZOVA = RGBColor(0xFF, 0x9A, 0x3C)
RUZOVA = RGBColor(0xF4, 0x72, 0xB6)
TYRKYS = RGBColor(0x34, 0xD3, 0x99)

NADPIS = "Segoe UI Black"
TEXT = "Segoe UI"
TEXT_TUCNY = "Segoe UI Semibold"

W, H = Inches(13.333), Inches(7.5)
PRYC = Inches(16)  # mimo plochu vpravo: odtud objekty přilétají


# ------------------------------------------------------------------ data
def cisla() -> dict:
    db = sqlite3.connect(KOREN / "data" / "kam.db")
    q = lambda sql: db.execute(sql).fetchone()[0]  # noqa: E731
    out = {
        "funkcnich": q("select count(*) from dataset where check_status = 'ok'"),
        "katalog": q("select count(*) from dataset"),
        "nacteno": q("select count(*) from dataset where loaded = 1"),
        "hodnot": q("select count(*) from observation"),
        "obci": q("select count(*) from territory where level = 'obec' and kraj_code = 86"),
        "zdroju": q("select count(distinct source_id) from dataset"),
        "roky": db.execute("select year, count(*) from observation group by year order by year").fetchall(),
    }
    db.close()
    return out


def cz(n: int) -> str:
    return f"{n:,}".replace(",", "\u00a0")


# ------------------------------------------------------------------ snímky aplikace
SNIMKY = {
    "obec": "obce/obec:574121", "scenar": "scenare?t=obec:574121&navrh=1", "pruvodce": "scenare",
    "stats": "statistiky", "mapa": "mapa", "prehled": "prehled", "hospodareni": "hospodareni?t=obec:574121",
}


def nafot() -> None:
    IMG.mkdir(parents=True, exist_ok=True)
    profil = VYSTUP / ".chrome"
    for jmeno, cesta in SNIMKY.items():
        cil = IMG / f"{jmeno}_cele.png"
        subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--hide-scrollbars", "--force-prefers-reduced-motion", "--force-dark-mode",
                        "--window-size=1600,1000", "--virtual-time-budget=15000", f"--user-data-dir={profil}", f"--screenshot={cil}",
                        f"http://localhost:8000/{cesta}"], capture_output=True, timeout=120)
        orez(cil, IMG / f"{jmeno}.png")
        cil.unlink(missing_ok=True)


def orez(zdroj: Path, cil: Path) -> None:
    """Bez postranního panelu a horní lišty: na snímku má být obsah, ne navigace."""
    im = Image.open(zdroj).convert("RGB")
    im.crop((250, 56, 1600, 1000)).save(cil)


# ------------------------------------------------------------------ stavební kameny
def pozadi(slide) -> None:
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = POZADI


def kruh(slide, jmeno, x, y, d, barva, pruhlednost=0):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
    s.name = jmeno
    s.fill.solid()
    s.fill.fore_color.rgb = barva
    s.line.fill.background()
    if pruhlednost:
        # průhlednost výplně (python-pptx ji neumí, píše se rovnou do XML)
        clr = s.fill._xPr.find(qn("a:solidFill"))[0]
        a = etree.SubElement(clr, qn("a:alpha"))
        a.set("val", str(int((100 - pruhlednost) * 1000)))
    return s


def text(slide, jmeno, x, y, w, h, behy, velikost, font=NADPIS, zarovnani=PP_ALIGN.LEFT, kotva=MSO_ANCHOR.TOP, radkovani=0.9):
    """`behy`: seznam (text, barva) — jeden řádek může mít dvě barvy."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tb.name = jmeno
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = kotva
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    radky: list[list[tuple[str, RGBColor]]] = [[]]
    for t, b in behy:
        casti = t.split("\n")
        for i, c in enumerate(casti):
            if i:
                radky.append([])
            if c:
                radky[-1].append((c, b))
    for i, r in enumerate(radky):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = zarovnani
        p.line_spacing = radkovani
        for t, b in r:
            run = p.add_run()
            run.text = t
            run.font.size = Pt(velikost)
            run.font.name = font
            run.font.color.rgb = b
    return tb


def obrazek(slide, jmeno, soubor, x, y, w, rotace=0.0):
    pic = slide.shapes.add_picture(str(IMG / f"{soubor}.png"), x, y, width=w)
    pic.name = jmeno
    pic.rotation = rotace
    pic.line.color.rgb = RGBColor(0x2A, 0x2E, 0x35)
    pic.line.width = Pt(1)
    return pic


def pruh(slide, x, y, w, barva):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Pt(6))
    s.name = "!!pruh"
    s.fill.solid()
    s.fill.fore_color.rgb = barva
    s.line.fill.background()
    return s


def graf(slide, roky, x, y, w, h, barva):
    """Plocha měřených hodnot po letech jako jeden tvar (kvůli Morphu)."""
    maxv = max(c for _, c in roky) or 1
    n = len(roky) - 1 or 1
    body = [(x + int(w * i / n), y + h - int(h * c / maxv)) for i, (_, c) in enumerate(roky)]
    fb = slide.shapes.build_freeform(x, y + h, scale=1.0)
    fb.add_line_segments(body + [(x + w, y + h)], close=True)
    s = fb.convert_to_shape()
    s.name = "!!graf"
    s.fill.solid()
    s.fill.fore_color.rgb = barva
    clr = s.fill._xPr.find(qn("a:solidFill"))[0]
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", "55000")
    s.line.color.rgb = barva
    s.line.width = Pt(3)
    return s


def morph(slide, prvni=False) -> None:
    """Přechod Morph (PowerPoint 2019+) s prolnutím pro starší verze."""
    sld = slide._element
    xml = (
        '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" '
        'xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main">'
        '<mc:Choice Requires="p159"><p:transition spd="slow" p14:dur="{dur}">{efekt}</p:transition></mc:Choice>'
        '<mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>'
    ).format(dur=1000 if prvni else 1600, efekt="<p:fade/>" if prvni else '<p159:morph option="byObject"/>')
    el = etree.fromstring(xml)
    # přechod patří za cSld/clrMapOvr a před timing
    timing = sld.find(qn("p:timing"))
    if timing is not None:
        timing.addprevious(el)
    else:
        sld.append(el)


# ------------------------------------------------------------------ snímky
def stavej(c: dict) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    prazdny = prs.slide_layouts[6]

    def snimek(nastaveni: dict, prvni=False):
        """Jeden snímek: všechny pojmenované objekty, ty nepoužité mimo plochu."""
        s = prs.slides.add_slide(prazdny)
        pozadi(s)
        k1 = nastaveni["kruh"]
        k2 = nastaveni["kruh2"]
        kruh(s, "!!kruh", *k1)
        kruh(s, "!!kruh2", *k2)
        g = nastaveni.get("graf")
        graf(s, c["roky"], *(g or (PRYC, Inches(4.2), Inches(13.3), Inches(3.3), FIALOVA)))
        for i, ob in enumerate(nastaveni.get("obrazky", []) + [None] * (3 - len(nastaveni.get("obrazky", [])))):
            if ob:
                obrazek(s, f"!!obr{i + 1}", *ob)
            else:
                obrazek(s, f"!!obr{i + 1}", ["obec", "scenar", "stats"][i], PRYC, Inches(1), Inches(6))
        st = nastaveni.get("stitek")
        if st:
            text(s, "!!stitek", Inches(0.8), Inches(0.6), Inches(6), Inches(0.5), [(st[0], st[1])], 16, TEXT_TUCNY)
            pruh(s, Inches(0.8), Inches(1.1), Inches(0.9), st[1])
        else:
            text(s, "!!stitek", PRYC, Inches(0.6), Inches(6), Inches(0.5), [("", BILA)], 16, TEXT_TUCNY)
            pruh(s, PRYC, Inches(1.1), Inches(0.9), BILA)
        v = nastaveni["velky"]
        text(s, "!!velky", *v["poloha"], v["behy"], v["velikost"])
        p = nastaveni.get("sub")
        if p:
            text(s, "!!sub", *p["poloha"], [(p["text"], SEDA)], p.get("velikost", 30), TEXT)
        for t in nastaveni.get("navic", []):
            text(s, *t)
        morph(s, prvni)
        return s

    # 1 — název týmu
    snimek({
        "kruh": (Inches(7.2), Inches(-1.6), Inches(8.6), ZELENA, 70),
        "kruh2": (Inches(1.0), Inches(6.2), Inches(0.9), ORANZOVA),
        "velky": {"poloha": (Inches(0.8), Inches(2.1), Inches(12), Inches(2.4)), "behy": [("Popojedem", BILA), (".", ZELENA)], "velikost": 140},
        "sub": {"poloha": (Inches(0.85), Inches(4.65), Inches(10), Inches(1.2)), "text": "KAM — Krajský analytický model\nHackathon Královéhradeckého kraje", "velikost": 28},
    }, prvni=True)

    # 2 — data: kolik
    snimek({
        "kruh": (Inches(-2.5), Inches(3.6), Inches(7.5), FIALOVA, 55),
        "kruh2": (Inches(11.6), Inches(0.7), Inches(1.0), ZELENA),
        "stitek": ("DATA", FIALOVA),
        "velky": {"poloha": (Inches(0.8), Inches(1.3), Inches(12.5), Inches(3.8)), "behy": [(cz(c["funkcnich"]), BILA)], "velikost": 230},
        "sub": {"poloha": (Inches(0.85), Inches(5.25), Inches(11), Inches(1.4)),
                "text": f"funkčních datových sad z {cz(c['zdroju'])} zdrojů\nověřených: v každé jsme našli skutečná data", "velikost": 30},
    })

    # 3 — data: jak je využíváme
    snimek({
        "kruh": (Inches(9.6), Inches(-2.2), Inches(5.4), FIALOVA),
        "kruh2": (Inches(0.9), Inches(6.6), Inches(0.5), ZELENA),
        "stitek": ("DATA", FIALOVA),
        "graf": (Inches(0), Inches(5.0), W, Inches(2.5), FIALOVA),
        "velky": {"poloha": (Inches(0.8), Inches(1.2), Inches(12.5), Inches(3.0)), "behy": [(cz(c["hodnot"]), BILA)], "velikost": 180},
        "sub": {"poloha": (Inches(0.85), Inches(4.25), Inches(11.5), Inches(0.8)),
                "text": f"hodnot počítáme z {c['nacteno']} sad · každá s původem", "velikost": 28},
    })

    # 4 — obce
    snimek({
        "kruh": (Inches(5.6), Inches(0.6), Inches(7.8), MODRA, 60),
        "kruh2": (Inches(5.6), Inches(0.55), Inches(0.5), MODRA),
        "stitek": ("OBCE", MODRA),
        "obrazky": [("obec", Inches(6.6), Inches(1.25), Inches(6.4), -2.5)],
        "velky": {"poloha": (Inches(0.8), Inches(1.3), Inches(6.2), Inches(4.4)), "behy": [(f"{c['obci']}\n", BILA), ("obcí", MODRA)], "velikost": 130},
        "sub": {"poloha": (Inches(0.85), Inches(5.75), Inches(5.6), Inches(1.4)), "text": "profil, rizika, rozpočet\na znak každé z nich", "velikost": 26},
    })

    # 5 — scénáře
    snimek({
        "kruh": (Inches(-3.0), Inches(-3.0), Inches(8.0), ORANZOVA, 40),
        "kruh2": (Inches(12.2), Inches(6.6), Inches(0.6), MODRA),
        "stitek": ("SCÉNÁŘE", ORANZOVA),
        "obrazky": [("scenar", Inches(7.6), Inches(2.3), Inches(5.4), 2.5), ("pruvodce", Inches(7.2), Inches(0.9), Inches(4.0), -4.0)],
        "velky": {"poloha": (Inches(0.8), Inches(1.5), Inches(6.5), Inches(4.0)), "behy": [("Co\nkdyž…?", BILA)], "velikost": 120},
        "sub": {"poloha": (Inches(0.85), Inches(5.6), Inches(5.8), Inches(1.4)), "text": "20 opatření a jejich dopad\nna školy, lékaře i daně", "velikost": 26},
    })

    # 6 — automatický návrh
    snimek({
        "kruh": (Inches(6.0), Inches(2.2), Inches(9.0), ORANZOVA, 55),
        "kruh2": (Inches(5.3), Inches(6.6), Inches(0.5), ZELENA),
        "stitek": ("SCÉNÁŘE", ORANZOVA),
        "obrazky": [("scenar", Inches(6.2), Inches(1.0), Inches(6.8), 0), ("pruvodce", PRYC, Inches(0.9), Inches(4.2), -4.0)],
        "velky": {"poloha": (Inches(0.8), Inches(1.6), Inches(5.6), Inches(3.2)), "behy": [("Navrhne\n", BILA), ("samo.", ORANZOVA)], "velikost": 95},
        "sub": {"poloha": (Inches(0.85), Inches(4.95), Inches(5.2), Inches(1.6)), "text": "najde problém obce\na spočítá, co a kolik přidat", "velikost": 26},
    })

    # 7 — grafy
    snimek({
        "kruh": (Inches(3.8), Inches(-4.2), Inches(9.0), RUZOVA, 45),
        "kruh2": (Inches(0.6), Inches(6.7), Inches(0.5), RUZOVA),
        "stitek": ("GRAFY", RUZOVA),
        "obrazky": [("stats", Inches(0.9), Inches(3.55), Inches(4.6), -3.0), ("mapa", Inches(4.6), Inches(3.25), Inches(4.6), 2.0), ("hospodareni", Inches(8.4), Inches(3.6), Inches(4.4), -1.5)],
        "velky": {"poloha": (Inches(0.8), Inches(1.3), Inches(12), Inches(1.9)), "behy": [("Data jsou ", BILA), ("vidět.", RUZOVA)], "velikost": 110},
        "sub": {"poloha": (Inches(8.0), Inches(0.62), Inches(4.6), Inches(0.6)), "text": "mapy · trendy · rozpočty · výhledy", "velikost": 18},
    })

    # 8 — API
    snimek({
        "kruh": (Inches(8.4), Inches(2.6), Inches(6.5), TYRKYS, 50),
        "kruh2": (Inches(11.8), Inches(0.6), Inches(0.8), ORANZOVA),
        "stitek": ("API", TYRKYS),
        "velky": {"poloha": (Inches(0.8), Inches(1.3), Inches(12), Inches(4.0)), "behy": [("Otevřené API\n", BILA), ("pro všechny.", TYRKYS)], "velikost": 120},
        "sub": {"poloha": (Inches(0.85), Inches(5.45), Inches(11), Inches(0.6)), "text": "/api/v1 · CSV · JSON · bez registrace", "velikost": 28},
        "navic": [("!!kod", Inches(0.85), Inches(6.2), Inches(11), Inches(0.6),
                   [("GET  hack-a-ton-hushkys.web.app/api/v1/hodnoty?uzemi=obec:574121", TYRKYS)], 20, "Consolas")],
    })

    # 9 — tým
    tym = [("Ondřej Hušek", "vývoj, architektura a AI"), ("Matěj Jirásek", "UX/UI design"),
           ("Josef Efler", "brand, marketing a legal"), ("Damián Roháček", "backend a testování")]
    snimek({
        "kruh": (Inches(9.2), Inches(-2.8), Inches(6.8), ZELENA, 70),
        "kruh2": (Inches(1.0), Inches(6.7), Inches(0.5), ORANZOVA),
        "stitek": ("TÝM", ZELENA),
        "velky": {"poloha": (Inches(0.8), Inches(1.35), Inches(12), Inches(2.4)), "behy": [("Popojedem", BILA), (".", ZELENA)], "velikost": 140},
        "navic": [(f"!!clen{i + 1}", Inches(0.85 + (i % 2) * 6.0), Inches(4.2 + (i // 2) * 1.4), Inches(5.6), Inches(1.2),
                   [(f"{j}\n", BILA), (r, SEDA)], 26, TEXT_TUCNY, PP_ALIGN.LEFT, MSO_ANCHOR.TOP, 1.05) for i, (j, r) in enumerate(tym)],
    })

    VYSTUP.mkdir(exist_ok=True)
    cil = VYSTUP / "KAM_Popojedem.pptx"
    prs.save(cil)
    return cil


if __name__ == "__main__":
    if "--snimky" in sys.argv:
        nafot()
    print(stavej(cisla()))
